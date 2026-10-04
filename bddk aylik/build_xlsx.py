import warnings; warnings.filterwarnings('ignore')
from calc import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import sys

OUT=sys.argv[1] if len(sys.argv)>1 else 'cikti.xlsx'
wb=Workbook(); HF=Font(name='Arial',bold=True,color='FFFFFF'); NF=Font(name='Arial',size=9); FILL=PatternFill('solid',fgColor='1C3044')

def sheet(name,df,cols,fmt='#,##0.0',first_w=14):
    ws=wb.create_sheet(name)
    ws.append(['Grup','Dönem']+cols)
    for c in ws[1]: c.font=HF; c.fill=FILL; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for g in GROUPS:
        sub=df[df.g==g].sort_values('d')
        for _,r in sub.iterrows():
            ws.append([GNAME[g],r.d.strftime('%Y-%m')]+[None if pd.isna(r[c]) else float(r[c]) for c in cols])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=NF
            if c.column>2: c.number_format=fmt
    ws.freeze_panes='C2'; ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=9
    for i in range(3,len(cols)+3): ws.column_dimensions[get_column_letter(i)].width=first_w
    ws.row_dimensions[1].height=42
    return ws

raw_cols=['nakTP','nakYP','mkTP','mkYP','krTP','krYP','tkTP','tkYP','fkTP','fkYP','ozelKars','akTP','akYP','mvTP','mvYP','vdsTP','vdsYP','topTP','topYP','sermBz','ozk','reeskTP','reeskYP','htmTP','htmYP','pasTP','pasYP',
 'faizGel','faizGid','ucrGel','ucrGid','ticari','digGel','personel','digIsl','karsilik','vergi','netKar','krFaiz','mkFaiz','digFaizGel','mvFaiz','topFaiz','digFaizGid','fkGel','fkGid','nakdiKom',
 'tuk','konut','tasit','ihtiyac','kk','tukKK','tkTuk','tkKonut','tkTasit','tkIhtiyac','tkKK','tkTukKK','kobiTP','kobiYP','kobi','tkKobi','tlVds','tlTot','dthVds','dthTot','kmVds','kmTot','anaSerm','katkiSerm','yasalOzk','cekirdek','rak','subeYI','subeYD','persYI','persYD','usd']
sheet('Ham Veri',D,raw_cols,'#,##0')
lvl=['aktif','kredi','canli','canliTP','canliYP','takip','fk','menkul','mevduat','vadeli','vadesiz','vadeliTP','vadeliYP','ypFon','km','tuzel','ticKur','ticKurTP','ticKurYP','tkTicKur','toptan','sube','pers','avgAktif','avgOzk','avgCanli','avgMev']+[c for c in D.columns if c.endswith('_12')]
sheet('Türetilmiş Tutarlar',D,lvl,'#,##0')
ratios=[c for c in D.columns if c.startswith('g_') or c.startswith('gy_')]+['aktifTPag','pasifTPag','fonTPag','aktifPay','tukCanli','konutCanli','tasitCanli','ihtiyacCanli','kkCanli','tuzelCanli','ticCanli','kobiCanli','fkCanli','npl','ozelKarsOran','nplTuk','nplKobi','nplTic','syr','cekSyr','anaSyr','roa','roe','nim','giderGelir','opexAktif','krediGetiri','fonMaliyet','tlKrediMev','ypKrediMev','krediMev','tlMkMev','ypMkMev','mkMev','mkTPak','mkYPak','mkAk','subePers','roaNetFaiz','roaUcret','roaTicari','roaDiger','roaOpex','roaKars','roaVergi','mKredi','mMenkul','mDigGel','mFon','mToptan','mDigGid','ucrGelAk','netFaizBrut','netFaizOpex','ucrBrut','ucrOpex','tlVdlYuk','ypVdlYuk','vdlYuk','tlVdsYuk','ypVdsYuk','vdsYuk']
D['aktifPay']=D.aktif/D.groupby('d').aktif.transform(lambda x: x[D.loc[x.index,'g']=='S'].iloc[0])*100
sheet('Rasyolar ve Büyümeler',D[D.d>='2023-01-31'],ratios,'0.00')

ws=wb.active; ws.title='Tanımlar'
defs=[('Kaynak','BDDK Aylık Bülten (Temel Gösterim) – tablolar: Bilanço(1), Kar Zarar(2), Tüketici Kredileri(4), KOBİ(6), Mevduat Vade(10), Sermaye Yeterliliği(12), Diğer Bilgiler(16); USD bilanço (kur için)'),
 ('Gruplar','S=Sektör, M=Mevduat, KA=Mevduat-Kamu, YO=Mevduat-Yerli Özel, YA=Mevduat-Yabancı, KB=Katılım'),
 ('Birim','Ham veri ve tutarlar Milyon TL; rasyolar %; g_ = yıllık büyüme (%), gy_ = yıl içi kümülatif kalemde yıllık değişim (%)'),
 ('nak','Nakit + TCMB + Para piyasası + Bankalar + Zorunlu karşılık + MK ödünç + Ters repo'),('mk','GUD K/Z + GUD OCI + İtfa edilmiş maliyet (HTM) menkul değerler'),
 ('kredi / canli','kredi = Krediler + Takip + Fin.Kiralama ; canli = Krediler + Fin.Kiralama'),('toptan','Para piyasasına borçlar + Bankalara borçlar + İhraç edilen MK'),
 ('km','Kıymetli maden depo hesapları (yurt içi + yurt dışı, Mevduat Vade tablosu); ypFon = YP mevduat − km'),
 ('faizGel','Toplam faiz gelirleri + nakdi kredi komisyonları ; ucrGel = gayrinakdi kredi komisyonları + bankacılık hizmet gelirleri'),
 ('ticari','Sermaye piyasası + kambiyo + net parasal pozisyon ; digGel = alınan kâr payları + aktif satış + diğer faiz dışı gelirler + olağanüstü'),
 ('personel / digIsl','personel = personel gid. + kıdem tazminatı ; digIsl = amortisman + vergi/resim/harç + diğer faiz dışı giderler'),
 ('karsilik','Takipteki alacaklar özel provizyonu + genel karşılık + MK değer azalma + iştirak değer azalma + diğer provizyonlar'),
 ('_12','12 aylık kümülatif: YTD(t) − YTD(t−12) + YTD(Aralık t−1)'),('avg*','13 aylık basit ortalama'),
 ('roa/roe/nim','netKar_12/avgAktif ; netKar_12/avgOzk ; netFaiz_12/avgAktif'),('krediGetiri','(krFaiz_12+fkGel_12)/avgCanli ; fonMaliyet = mvFaiz_12/avgMev'),
 ('npl','takip/kredi ; nplTuk = tkTukKK/(tukKK+tkTukKK) ; nplKobi = tkKobi/(kobi+tkKobi) ; nplTic = kalan takip / (ticKur+kalan takip)'),
 ('syr','yasalOzk/rak ; cekSyr = cekirdek/rak ; anaSyr = anaSerm/rak'),('mkGr','Menkul kıymet büyümesi: HTM hariç, MK reeskontu dahil'),
 ('*_usd','TL tutar / örtük USD-TL kuru (Sektör toplam aktif TL / USD)'),('Not','ekordion.com.tr Bankacılık Monitörü veri dosyası (fetch_bddk.py → calc.py → export_web.py); hücreler değer olarak yazılmıştır.')]
ws.append(['Alan','Tanım'])
for c in ws[1]: c.font=HF; c.fill=FILL
for k,v in defs: ws.append([k,v])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font=NF; c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=20; ws.column_dimensions['B'].width=140
wb.save(OUT); print('saved',OUT)
